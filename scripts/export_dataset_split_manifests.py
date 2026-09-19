"""Export ISIC and Polyp train/validation/test manifests into one workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_ROOT = PROJECT_ROOT / "Data"
OUTPUT_PATH = DATA_ROOT / "dataset_split_manifests.xlsx"
SPLITS = ("train", "validation", "test")
POLYP_DATASETS = ("Kvasir", "CVC-ClinicDB", "CVC-ColonDB", "CVC-300", "ETIS-LaribPolypDB")


def read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return [], []
    return rows[0], rows[1:]


def read_manifest(source: Path) -> tuple[list[str], list[list[str]]]:
    return read_csv(source)


def write_sheet(
    workbook: Workbook,
    title: str,
    source: Path,
    headers: list[str] | None = None,
    rows: list[list[str]] | None = None,
) -> int:
    if headers is None or rows is None:
        headers, rows = read_manifest(source)
    worksheet = workbook.create_sheet(title)
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column_index, header in enumerate(headers, start=1):
        longest = max([len(header), *(len(row[column_index - 1]) for row in rows if len(row) >= column_index)])
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 12), 70)
    return len(rows)


def main() -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Dataset", "Split", "Samples", "Source CSV", "Worksheet"])

    records: list[tuple[str, str, Path, str]] = []
    isic_root = DATA_ROOT / "ISIC2018" / "manifests"
    for split in SPLITS:
        records.append(("ISIC2018", split, isic_root / f"{split}.csv", f"ISIC_{split}"))

    polyp_root = DATA_ROOT / "kvasir-seg" / "organized"
    for dataset in POLYP_DATASETS:
        for split in SPLITS:
            records.append((dataset, split, polyp_root / dataset / f"{split}.csv", f"{dataset}_{split}"))

    manifest_worksheets = 0
    for dataset, split, source, worksheet_name in records:
        headers, rows = read_manifest(source)
        if not rows:
            continue
        count = write_sheet(workbook, worksheet_name, source, headers, rows)
        summary.append((dataset, split, count, str(source.relative_to(PROJECT_ROOT)), worksheet_name))
        manifest_worksheets += 1

    write_sheet(workbook, "Polyp_Split_Summary", polyp_root / "split_summary.csv")
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column, width in zip(("A", "B", "C", "D", "E"), (22, 14, 12, 58, 32)):
        summary.column_dimensions[column].width = width

    workbook.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH} with {manifest_worksheets} non-empty manifest worksheets.")


if __name__ == "__main__":
    main()

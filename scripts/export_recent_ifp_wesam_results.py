#!/usr/bin/env python3
"""Export the recent IFP-WeSAM experiments into one auditable workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output_current" / "recent_ifp_wesam_experiments.xlsx"

ISIC_EXPERIMENTS = (
    ("ISIC 1% IFP semi-supervised", "1%", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "FP32", ROOT / "output_current/isic/1pct"),
    ("ISIC 10% IFP semi-supervised", "10%", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "FP32", ROOT / "output_current/isic/10pct"),
    ("ISIC exact 10 GT IFP semi-supervised", "10 images", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "BF16", ROOT / "output_current/isic/10gt_ifp_teacher_student_retry5_bf16_no_anchor"),
)

POLYP_EXPERIMENTS = (
    ("Polyp 1% IFP semi-supervised", "1%", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "FP32", ROOT / "output_current/Polyp/1pct"),
    ("Polyp 10% IFP semi-supervised", "10%", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "FP32", ROOT / "output_current/Polyp/10pct"),
    ("Polyp exact 5 GT IFP semi-supervised", "5 images", "GT interior", "IFP top-1", "EMA Teacher pseudo-mask", "IFP top-1", "FP32", ROOT / "output_current/Polyp/5gt_ifp_teacher_student"),
    ("Polyp 1% DINO-only", "1%", "GT interior", "DINO top-1", "EMA Teacher pseudo-mask", "DINO top-1", "FP32", ROOT / "output_current/Polyp/1pct_dino_only"),
    ("Polyp 1% CLIP-only", "1%", "GT interior", "CLIP top-1", "EMA Teacher pseudo-mask", "CLIP top-1", "BF16", ROOT / "output_current/Polyp/1pct_clip_only_bf16_retry3"),
    ("Polyp 1% No-Prompt", "1%", "None", "None", "EMA Teacher pseudo-mask", "None", "FP32", ROOT / "output_current/Polyp/1pct_no_prompt_teacher_student"),
    ("Polyp Full-GT + IFP Prompt", "100%", "IFP top-1", "n.a.", "None (all GT)", "IFP top-1", "FP32", ROOT / "output_current/Polyp/full_gt_ifp_prompt"),
    ("Polyp Full-GT + GT Prompt (oracle)", "100%", "GT interior", "n.a.", "None (all GT)", "GT interior", "FP32", ROOT / "output_current/Polyp/full_gt_gt_prompt"),
)


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise FileNotFoundError(path)
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def style(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        width = max(len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 48)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"


def best_validation(rows: list[dict[str, str]]) -> dict[str, str]:
    candidates = [row for row in rows if row["Name"].endswith("_student")]
    if not candidates:
        raise ValueError("No Student validation records found")
    return max(candidates, key=lambda row: float(row["Mean IoU"]))


def isic_final_test(rows: list[dict[str, str]]) -> dict[str, str]:
    candidates = [row for row in rows if row["Name"].endswith("_final_test_best_student")]
    if len(candidates) != 1:
        raise ValueError("Expected exactly one final ISIC test record")
    return candidates[0]


def main() -> None:
    book = Workbook()
    isic = book.active
    isic.title = "ISIC Results"
    isic.append((
        "Experiment", "GT budget", "Labeled training point", "Unlabeled training point",
        "Unlabeled supervision", "Validation / test point", "Training precision",
        "Best validation epoch", "Best validation mIoU", "Best validation F1",
        "Test images", "Test mIoU", "Test F1", "Result source",
    ))
    for name, budget, labeled_point, unlabeled_point, unlabeled_supervision, eval_point, precision, directory in ISIC_EXPERIMENTS:
        source = directory / "ISIC" / "metrics.csv"
        rows = read_csv(source)
        best = best_validation(rows)
        final = isic_final_test(rows)
        isic.append((
            name, budget, labeled_point, unlabeled_point, unlabeled_supervision, eval_point, precision,
            int(best["epoch"]), float(best["Mean IoU"]), float(best["Mean F1"]), 1000,
            float(final["Mean IoU"]), float(final["Mean F1"]), str(source),
        ))

    polyp = book.create_sheet("Polyp 5 Test Sets")
    polyp.append((
        "Experiment", "GT budget", "Labeled training point", "Unlabeled training point",
        "Unlabeled supervision", "Test point", "Training precision", "Test set", "Images",
        "Best validation epoch", "Best validation mIoU", "Best validation F1", "Test mIoU", "Test F1", "Result source",
    ))
    for name, budget, labeled_point, unlabeled_point, unlabeled_supervision, eval_point, precision, directory in POLYP_EXPERIMENTS:
        validation_source = directory / "Polyp" / "metrics.csv"
        test_source = directory / "test_metrics.csv"
        validation_rows = read_csv(validation_source)
        test_rows = read_csv(test_source)
        best = best_validation(validation_rows)
        for row in test_rows:
            polyp.append((
                name, budget, labeled_point, unlabeled_point, unlabeled_supervision, eval_point, precision,
                row["Dataset"], int(row["Images"]), int(row.get("Best validation epoch", best["epoch"])),
                float(best["Mean IoU"]), float(best["Mean F1"]), float(row["Mean IoU"]),
                float(row["Mean F1"]), str(test_source),
            ))

    for sheet in book.worksheets:
        style(sheet)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)

    verified = load_workbook(OUTPUT, read_only=True)
    expected = {"ISIC Results", "Polyp 5 Test Sets"}
    if set(verified.sheetnames) != expected:
        raise RuntimeError(f"Unexpected sheets: {verified.sheetnames}")
    print(OUTPUT)


if __name__ == "__main__":
    main()

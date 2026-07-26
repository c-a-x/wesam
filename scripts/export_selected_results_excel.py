#!/usr/bin/env python3
"""Combine selected ISIC and Polyp 1%/10% CSV metrics into one workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output_current/ISIC_Polyp_1pct_10pct_results.xlsx"
ISIC_PATHS = {
    "1%": ROOT / "output_current/isic/1pct/ISIC/metrics.csv",
    "10%": ROOT / "output_current/isic/10pct/ISIC/metrics.csv",
}
POLYP_PATHS = {
    "1%": ROOT / "output_current/Polyp/1pct/test_metrics.csv",
    "10%": ROOT / "output_current/Polyp/10pct/test_metrics.csv",
}


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 44)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"


def weighted_polyp(rows: list[dict[str, str]]) -> tuple[int, float, float]:
    images = sum(int(row["Images"]) for row in rows)
    mean_iou = sum(int(row["Images"]) * float(row["Mean IoU"]) for row in rows) / images
    mean_f1 = sum(int(row["Images"]) * float(row["Mean F1"]) for row in rows) / images
    return images, mean_iou, mean_f1


def main() -> None:
    isic = {ratio: read_csv(path) for ratio, path in ISIC_PATHS.items()}
    polyp = {ratio: read_csv(path) for ratio, path in POLYP_PATHS.items()}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append((
        "Domain", "Labeled Ratio", "Evaluation", "Dataset", "Images",
        "Epoch", "Mean IoU", "Mean F1", "Source CSV",
    ))

    for ratio, rows in isic.items():
        baseline = next(row for row in rows if row["epoch"] == "0")
        validation_rows = [row for row in rows if row["Name"].endswith("_student")]
        best = max(validation_rows, key=lambda row: float(row["Mean IoU"]))
        final = next(row for row in rows if row["Name"].endswith("_final_test_best_student"))
        for label, row in (("Validation baseline", baseline), ("Best validation", best), ("Final test", final)):
            summary.append((
                "ISIC", ratio, label, "ISIC", None, int(row["epoch"]),
                float(row["Mean IoU"]), float(row["Mean F1"]), str(ISIC_PATHS[ratio]),
            ))

    for ratio, rows in polyp.items():
        for row in rows:
            summary.append((
                "Polyp", ratio, "Independent test", row["Dataset"], int(row["Images"]),
                int(row.get("Best validation epoch", "0") or 0), float(row["Mean IoU"]),
                float(row["Mean F1"]), str(POLYP_PATHS[ratio]),
            ))
        images, mean_iou, mean_f1 = weighted_polyp(rows)
        summary.append((
            "Polyp", ratio, "Image-weighted test aggregate", "All five test sets", images,
            None, mean_iou, mean_f1, str(POLYP_PATHS[ratio]),
        ))

    isic_sheet = workbook.create_sheet("ISIC_All_Epochs")
    isic_sheet.append(("Labeled Ratio", "Dataset", "Name", "Prompt", "Mean IoU", "Mean F1", "Epoch", "Record Type"))
    for ratio, rows in isic.items():
        for row in rows:
            record_type = (
                "Final test" if row["Name"].endswith("_final_test_best_student")
                else "Validation baseline" if row["epoch"] == "0"
                else "Validation"
            )
            isic_sheet.append((
                ratio, row["Dataset"], row["Name"], row["Prompt"], float(row["Mean IoU"]),
                float(row["Mean F1"]), int(row["epoch"]), record_type,
            ))

    polyp_sheet = workbook.create_sheet("Polyp_Test_Sets")
    polyp_sheet.append(("Labeled Ratio", "Dataset", "Images", "Mean IoU", "Mean F1", "Best Validation Epoch", "Source CSV"))
    for ratio, rows in polyp.items():
        for row in rows:
            polyp_sheet.append((
                ratio, row["Dataset"], int(row["Images"]), float(row["Mean IoU"]),
                float(row["Mean F1"]), row.get("Best validation epoch", ""), str(POLYP_PATHS[ratio]),
            ))

    comparison = workbook.create_sheet("Comparison")
    comparison.append(("Domain", "Evaluation / Dataset", "1% Mean IoU", "10% Mean IoU", "10%-1% IoU", "1% Mean F1", "10% Mean F1", "10%-1% F1"))
    for label, selector in (
        ("Best validation", lambda rows: max((row for row in rows if row["Name"].endswith("_student")), key=lambda row: float(row["Mean IoU"]))),
        ("Final test", lambda rows: next(row for row in rows if row["Name"].endswith("_final_test_best_student"))),
    ):
        one, ten = selector(isic["1%"]), selector(isic["10%"])
        comparison.append((
            "ISIC", label, float(one["Mean IoU"]), float(ten["Mean IoU"]), float(ten["Mean IoU"]) - float(one["Mean IoU"]),
            float(one["Mean F1"]), float(ten["Mean F1"]), float(ten["Mean F1"]) - float(one["Mean F1"]),
        ))
    for dataset in [row["Dataset"] for row in polyp["1%"]]:
        one = next(row for row in polyp["1%"] if row["Dataset"] == dataset)
        ten = next(row for row in polyp["10%"] if row["Dataset"] == dataset)
        comparison.append((
            "Polyp", dataset, float(one["Mean IoU"]), float(ten["Mean IoU"]), float(ten["Mean IoU"]) - float(one["Mean IoU"]),
            float(one["Mean F1"]), float(ten["Mean F1"]), float(ten["Mean F1"]) - float(one["Mean F1"]),
        ))
    one_images, one_iou, one_f1 = weighted_polyp(polyp["1%"])
    ten_images, ten_iou, ten_f1 = weighted_polyp(polyp["10%"])
    if one_images != ten_images:
        raise ValueError("Polyp 1% and 10% test image counts differ")
    comparison.append((
        "Polyp", "All five test sets (image-weighted)", one_iou, ten_iou, ten_iou - one_iou,
        one_f1, ten_f1, ten_f1 - one_f1,
    ))

    for sheet in workbook.worksheets:
        style_sheet(sheet)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)

    # Verify the written workbook can be opened and contains the expected sheets.
    verified = load_workbook(OUTPUT, read_only=True)
    expected = {"Summary", "ISIC_All_Epochs", "Polyp_Test_Sets", "Comparison"}
    if set(verified.sheetnames) != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {verified.sheetnames}")
    print(OUTPUT)


if __name__ == "__main__":
    main()
